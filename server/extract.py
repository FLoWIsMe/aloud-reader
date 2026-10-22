"""Text extraction: files, pasted text and web pages become a flat list of blocks."""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from bs4 import BeautifulSoup, Tag
from markdown_it import MarkdownIt

from .segment import Block, normalize_whitespace

TEXT_EXTENSIONS = {".txt", ".text", ".log"}
MARKDOWN_EXTENSIONS = {".md", ".markdown", ".mdown", ".mkd"}
HTML_EXTENSIONS = {".html", ".htm", ".xhtml"}
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS | MARKDOWN_EXTENSIONS | HTML_EXTENSIONS | {".pdf", ".docx", ".epub", ".rtf"}
)

_BLOCK_TAGS = [
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "p",
    "li",
    "blockquote",
    "pre",
    "figcaption",
    "dt",
    "dd",
    "td",
    "th",
    "caption",
    "summary",
]
_SKIP_TAGS = {
    "script",
    "style",
    "noscript",
    "template",
    "svg",
    "canvas",
    "nav",
    "header",
    "footer",
    "aside",
    "form",
    "button",
    "iframe",
    "video",
    "audio",
    "picture",
}
_MARKDOWN_SIGNALS = re.compile(r"^(#{1,6}\s|\s*[-*+]\s|\s*\d+[.)]\s|>\s|```)", re.MULTILINE)
_PAGE_NUMBER = re.compile(r"^\s*(?:page\s*)?\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?\s*$", re.IGNORECASE)
_HYPHEN_BREAK = re.compile(r"(\w)-\n(?=[a-z])")


class ExtractionError(ValueError):
    """The input could not be turned into readable text."""


@dataclass
class Extracted:
    blocks: list[Block] = field(default_factory=list)
    title: str | None = None

    def has_text(self) -> bool:
        return any(normalize_whitespace(block.text) for block in self.blocks)


# Plain text and Markdown


def extract_plain_text(text: str) -> Extracted:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks: list[Block] = []
    for raw in re.split(r"\n\s*\n", text):
        paragraph = normalize_whitespace(raw.replace("\n", " "))
        if paragraph:
            blocks.append(Block("paragraph", paragraph))
    return Extracted(blocks)


def extract_markdown(text: str) -> Extracted:
    html = MarkdownIt("commonmark", {"breaks": False}).enable("table").render(text)
    extracted = blocks_from_html(html)
    if not extracted.has_text():
        return extract_plain_text(text)
    return extracted


def extract_pasted_text(text: str) -> Extracted:
    stripped = text.strip()
    if stripped.startswith("<") and "</" in stripped:
        return extract_html(stripped)
    if _MARKDOWN_SIGNALS.search(stripped):
        return extract_markdown(stripped)
    return extract_plain_text(stripped)


# HTML


def _element_text(element: Tag) -> str:
    for br in element.find_all("br"):
        br.replace_with(" ")
    return normalize_whitespace(element.get_text(""))


_TITLE_SEPARATORS = (" | ", " - ", " \u2013 ", " \u2014 ", " \u00b7 ", " :: ")


def clean_title(title: str) -> str:
    """Drop the site name that pages append to their titles."""
    title = normalize_whitespace(title)
    for separator in _TITLE_SEPARATORS:
        if separator in title:
            parts = [part.strip() for part in title.split(separator) if part.strip()]
            if parts:
                first = parts[0]
                return first if len(first) >= 8 or len(first.split()) >= 2 else max(parts, key=len)
    return title


def _title_for(extracted: Extracted, fallback: str | None) -> str | None:
    for block in extracted.blocks[:3]:
        if block.kind == "heading" and block.level == 1 and 2 <= len(block.text) <= 200:
            return block.text
    return clean_title(fallback) if fallback else None


def blocks_from_html(html: str) -> Extracted:
    """Walk block-level elements in document order, keeping headings as headings."""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup.find_all(list(_SKIP_TAGS)):
        tag.decompose()
    title_tag = soup.find("title")
    title = normalize_whitespace(title_tag.get_text()) if title_tag else None
    root = soup.body or soup
    blocks: list[Block] = []
    for element in root.find_all(_BLOCK_TAGS):
        if element.find(_BLOCK_TAGS):
            continue  # a container; its children will be visited on their own
        text = _element_text(element)
        if not text:
            continue
        name = element.name.lower()
        if name[0] == "h" and name[1:].isdigit():
            blocks.append(Block("heading", text, int(name[1])))
        else:
            blocks.append(Block("paragraph", text))
    if not blocks:
        for line in root.get_text("\n").split("\n"):
            text = normalize_whitespace(line)
            if text:
                blocks.append(Block("paragraph", text))
    return Extracted(blocks, title or None)


def extract_html(html: str, url: str | None = None) -> Extracted:
    """Pull the article out of a page with trafilatura, falling back to a plain walk."""
    try:
        import trafilatura
    except ImportError:  # pragma: no cover - trafilatura is a declared dependency
        return blocks_from_html(html)
    document = trafilatura.bare_extraction(
        html,
        url=url,
        output_format="markdown",
        include_formatting=True,
        include_links=False,
        include_tables=True,
        include_images=False,
        favor_recall=True,
        with_metadata=True,
    )
    markdown = getattr(document, "text", None) if document else None
    if markdown and normalize_whitespace(markdown):
        extracted = extract_markdown(markdown)
        extracted.title = _title_for(extracted, getattr(document, "title", None))
        if extracted.has_text():
            return extracted
    fallback = blocks_from_html(html)
    fallback.title = _title_for(fallback, fallback.title)
    return fallback


# PDF


def reflow_pages(pages: list[str]) -> list[Block]:
    """Rebuild paragraphs from the line-oriented text PDF extractors produce."""
    pages = [page.replace("\r\n", "\n").replace("\r", "\n") for page in pages]
    # Lines repeated across many pages are running headers or footers.
    line_pages: Counter[str] = Counter()
    for page in pages:
        for line in {normalize_whitespace(line) for line in page.split("\n")}:
            if line:
                line_pages[line] += 1
    repeated = {
        line
        for line, count in line_pages.items()
        if count >= 3 and count >= 0.4 * len(pages) and len(line) < 120
    }

    blocks: list[Block] = []
    current: list[str] = []

    def flush() -> None:
        if current:
            text = normalize_whitespace(" ".join(current))
            if text:
                blocks.append(Block("paragraph", text))
            current.clear()

    for page in pages:
        page = _HYPHEN_BREAK.sub(r"\1", page)
        lines = [normalize_whitespace(line) for line in page.split("\n")]
        lengths = sorted(len(line) for line in lines if line)
        median = lengths[len(lengths) // 2] if lengths else 0
        for line in lines:
            if not line:
                flush()
                continue
            if line in repeated or _PAGE_NUMBER.match(line):
                continue
            current.append(line)
            ends_sentence = line[-1] in ".!?:\u201d\"'"
            short = median and len(line) < 0.6 * median
            if ends_sentence and short:
                flush()
    flush()
    return _promote_headings(blocks)


def _promote_headings(blocks: list[Block]) -> list[Block]:
    """Short standalone lines without terminal punctuation read as section titles."""
    for index, block in enumerate(blocks):
        text = block.text
        if (
            index + 1 < len(blocks)
            and 2 <= len(text) <= 70
            and text[-1] not in ".!?,;:"
            and not text[0].islower()
            and len(text.split()) <= 10
        ):
            block.kind = "heading"
            block.level = 2
    return blocks


def extract_pdf(data: bytes) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover
        raise ExtractionError("PDF support is not installed") from error
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as error:
                raise ExtractionError("This PDF is password protected") from error
        pages = [page.extract_text() or "" for page in reader.pages]
    except ExtractionError:
        raise
    except Exception as error:
        raise ExtractionError(f"Could not read this PDF: {error}") from error
    title = None
    try:
        if reader.metadata and reader.metadata.title:
            title = normalize_whitespace(str(reader.metadata.title)) or None
    except Exception:
        title = None
    blocks = reflow_pages(pages)
    if not any(block.text for block in blocks):
        raise ExtractionError(
            "No text found in this PDF. It may be scanned images, which need OCR first."
        )
    return Extracted(blocks, title)


# DOCX


def extract_docx(data: bytes) -> Extracted:
    try:
        import docx
    except ImportError as error:  # pragma: no cover
        raise ExtractionError("Word support is not installed") from error
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as error:
        raise ExtractionError(f"Could not read this Word document: {error}") from error
    blocks: list[Block] = []
    for paragraph in document.paragraphs:
        text = normalize_whitespace(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        match = re.match(r"(?:Heading|Title)\s*(\d)?", style, re.IGNORECASE)
        if match:
            level = int(match.group(1)) if match.group(1) else 1
            blocks.append(Block("heading", text, level))
        else:
            blocks.append(Block("paragraph", text))
    for table in document.tables:
        for row in table.rows:
            cells = [normalize_whitespace(cell.text) for cell in row.cells]
            text = ". ".join(cell for cell in cells if cell)
            if text:
                blocks.append(Block("paragraph", text))
    title = None
    try:
        title = normalize_whitespace(document.core_properties.title or "") or None
    except Exception:
        title = None
    return Extracted(blocks, title)


# EPUB


def extract_epub(data: bytes) -> Extracted:
    try:
        import ebooklib
        from ebooklib import epub
    except ImportError as error:  # pragma: no cover
        raise ExtractionError("EPUB support is not installed") from error
    import logging
    import tempfile

    logging.getLogger("ebooklib").setLevel(logging.ERROR)
    with tempfile.NamedTemporaryFile(suffix=".epub") as handle:
        handle.write(data)
        handle.flush()
        try:
            book = epub.read_epub(handle.name, options={"ignore_ncx": True})
        except Exception as error:
            raise ExtractionError(f"Could not read this EPUB: {error}") from error
    title = None
    try:
        titles = book.get_metadata("DC", "title")
        if titles:
            title = normalize_whitespace(titles[0][0]) or None
    except Exception:
        title = None
    items = {item.get_id(): item for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT)}
    ordered = []
    for entry in book.spine:
        item_id = entry[0] if isinstance(entry, tuple) else entry
        if item_id in items:
            ordered.append(items.pop(item_id))
    ordered.extend(items.values())
    blocks: list[Block] = []
    for item in ordered:
        try:
            content = item.get_content().decode("utf-8", errors="replace")
        except Exception:
            continue
        blocks.extend(blocks_from_html(content).blocks)
    return Extracted(blocks, title)


# RTF (plain stripping is enough for reading aloud)


def extract_rtf(data: bytes) -> Extracted:
    text = data.decode("latin-1", errors="replace")
    text = re.sub(r"\\par[d]?\b", "\n\n", text)
    text = re.sub(r"\\'[0-9a-f]{2}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = re.sub(r"[{}]", "", text)
    return extract_plain_text(text)


# Dispatch


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def sniff_extension(filename: str, data: bytes) -> str:
    extension = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if extension in SUPPORTED_EXTENSIONS:
        return extension
    if data.startswith(b"%PDF"):
        return ".pdf"
    if data.startswith(b"{\\rtf"):
        return ".rtf"
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            names = set()
        if "word/document.xml" in names:
            return ".docx"
        if "META-INF/container.xml" in names or "mimetype" in names:
            return ".epub"
    head = data[:2048].lstrip().lower()
    if head.startswith(b"<!doctype html") or head.startswith(b"<html"):
        return ".html"
    return ".txt"


def extract_file(filename: str, data: bytes) -> Extracted:
    extension = sniff_extension(filename, data)
    if extension == ".pdf":
        extracted = extract_pdf(data)
    elif extension == ".docx":
        extracted = extract_docx(data)
    elif extension == ".epub":
        extracted = extract_epub(data)
    elif extension == ".rtf":
        extracted = extract_rtf(data)
    elif extension in HTML_EXTENSIONS:
        extracted = extract_html(decode_text(data))
    elif extension in MARKDOWN_EXTENSIONS:
        extracted = extract_markdown(decode_text(data))
    else:
        extracted = extract_plain_text(decode_text(data))
    if not extracted.has_text():
        raise ExtractionError("No readable text was found in this file")
    return extracted


def title_from_filename(filename: str) -> str:
    stem = PurePosixPath(filename.replace("\\", "/")).stem
    stem = re.sub(r"[_\-]+", " ", stem)
    return normalize_whitespace(stem) or "Untitled"


def title_from_blocks(blocks: list[Block]) -> str:
    for block in blocks:
        if block.kind == "heading" and block.text:
            return block.text[:120]
    for block in blocks:
        if block.text:
            words = block.text.split()
            title = " ".join(words[:12])
            return title + ("\u2026" if len(words) > 12 else "")
    return "Untitled"
